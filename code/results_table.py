"""Persist experiment results and populate the comparison workbook.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
import math
from pathlib import Path

EXPERIMENT_COLUMNS = ("exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay", "batch", "epochs", "hidden", "dropout", "clip_norm", "precision", "init", "seed", "step0_loss", "best_val_loss", "best_epoch", "final_train_loss", "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB", "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes")

def _jsonable(value):
    if isinstance(value, dict): return {str(k): _jsonable(v) for k, v in value.items() if k != "best_state"}
    if isinstance(value, (list, tuple)): return [_jsonable(v) for v in value]
    if hasattr(value, "item"): return _jsonable(value.item())
    if isinstance(value, float) and not math.isfinite(value): return None
    return value


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    cfg = result.get("cfg", {}); exp_id = cfg.get("exp_id")
    if not exp_id: raise ValueError("result['cfg']['exp_id'] bị thiếu")
    out_dir = Path(results_dir); out_dir.mkdir(parents=True, exist_ok=True)
    payload = {"cfg": result.get("cfg", {}), "history": result.get("history", {}), "summary": result.get("summary", {})}
    path = out_dir / f"{exp_id}.json"; path.write_text(json.dumps(_jsonable(payload), ensure_ascii=False, indent=2), encoding="utf-8"); return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    directory = Path(results_dir)
    if not directory.exists(): return []
    return [json.loads(path.read_text(encoding="utf-8")) for path in sorted(directory.glob("*.json"))]


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg = result.get("cfg", {}); summary = result.get("summary", {}); scores = eval_scores or {}; opt = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}
    row = {"exp_id": cfg.get("exp_id"), "group": cfg.get("group"), "description": cfg.get("description"), "loss": str(cfg.get("loss", "")).upper(), "optimizer": opt.get(cfg.get("optimizer"), cfg.get("optimizer")), "lr": cfg.get("lr"), "weight_decay": cfg.get("weight_decay"), "batch": cfg.get("batch"), "epochs": cfg.get("epochs"), "hidden": "-".join(str(x) for x in cfg.get("hidden", ())), "dropout": cfg.get("dropout"), "clip_norm": "none" if cfg.get("clip_norm") is None else cfg.get("clip_norm"), "precision": cfg.get("precision"), "init": cfg.get("init"), "seed": cfg.get("seed"), "step0_loss": summary.get("step0_loss"), "best_val_loss": summary.get("best_val_loss"), "best_epoch": summary.get("best_epoch"), "final_train_loss": summary.get("final_train_loss"), "final_val_loss": summary.get("final_val_loss"), "val_acc": summary.get("val_acc"), "val_macro_f1": summary.get("val_macro_f1"), "time_per_epoch_s": summary.get("time_per_epoch_s"), "peak_mem_MB": summary.get("peak_mem_MB"), "diverged": summary.get("diverged"), "eval_acc": scores.get("eval_acc", scores.get("accuracy", scores.get("acc"))), "eval_macro_f1": scores.get("eval_macro_f1", scores.get("macro_f1")), "figure_file": f"figures/{cfg.get('exp_id')}.png", "notes": notes}
    return {key: row.get(key) for key in EXPERIMENT_COLUMNS}


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    from openpyxl import load_workbook
    workbook = load_workbook(template_path); worksheet = workbook["Experiments"]
    headers = {worksheet.cell(1, c).value: c for c in range(1, worksheet.max_column + 1)}
    formula_cols = {c for c in range(1, worksheet.max_column + 1) if isinstance(worksheet.cell(2, c).value, str) and worksheet.cell(2, c).value.startswith("=")}
    for r in range(2, max(worksheet.max_row, len(rows) + 1) + 1):
        row = rows[r - 2] if r - 2 < len(rows) else {}
        for key, col in headers.items():
            if col not in formula_cols: worksheet.cell(r, col).value = row.get(key)
    output = Path(out_path); output.parent.mkdir(parents=True, exist_ok=True); workbook.save(output)
